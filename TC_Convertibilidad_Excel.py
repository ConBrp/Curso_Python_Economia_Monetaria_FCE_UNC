# Importación de librerías.
import pandas as pd
import json
import requests
import openpyxl
from openpyxl.chart import LineChart, Reference, Series


# Descarga de cotizaciones dólar blue.
response = requests.get("https://api.bluelytics.com.ar/v2/evolution.json")
data = json.loads(response.text)
blue = pd.DataFrame(data)
blue = blue[blue["source"] == "Blue"]

blue = blue.drop("source", axis=1)

blue["Precio"] = (blue["value_sell"] + blue["value_buy"]) / 2
blue["date"] = pd.to_datetime(blue["date"])
blue.columns = ["Fecha", "Venta", "Compra", "Promedio"]

df = pd.read_excel("https://www.bcra.gob.ar/Pdfs/PublicacionesEstadisticas/series.xlsm", header=[0, 1, 2, 3, 4, 5, 6, 7, 8],
				   sheet_name=["RESERVAS", "BASE MONETARIA", "INSTRUMENTOS DEL BCRA"])
df["RESERVAS"].columns = [str(i) for i in range(1, len(df["RESERVAS"].columns) + 1)]
reservas = df["RESERVAS"][df["RESERVAS"]["17"] == "D"][["1", "3", "16"]].copy().reset_index(drop=1)
reservas.columns = ["Fecha", "Reservas", "TCOficial"]

df["BASE MONETARIA"].columns = [str(i) for i in range(1, len(df["BASE MONETARIA"].columns) + 1)]
base = df["BASE MONETARIA"][df["BASE MONETARIA"]["32"] == "D"][["1", "29"]].copy().reset_index(drop=1)
base.columns = ["Fecha", "Base Monetaria"]

df["INSTRUMENTOS DEL BCRA"].columns = [str(i) for i in range(1, len(df["INSTRUMENTOS DEL BCRA"].columns) + 1)]
pasivos = df["INSTRUMENTOS DEL BCRA"][["1", "2", "4", "5", "6", "7", "8"]].copy()
pasivos.columns = ["Fecha", "PasesPasivos", "PasesActivos", "Leliq", "Lebac", "NOCOM", "LevidUSD"]

dolar = pd.merge(pd.merge(base, reservas, on="Fecha", how="inner"), pasivos, on="Fecha", how="inner")

dolar["LevidPesos"] = dolar["LevidUSD"] * dolar["TCOficial"]
dolar["PasesPasivos"] = pd.to_numeric(dolar["PasesPasivos"], errors='coerce')
dolar["PasesActivos"] = pd.to_numeric(dolar["PasesActivos"], errors='coerce')
dolar["Leliq"] = pd.to_numeric(dolar["Leliq"], errors='coerce')
dolar["Lebac"] = pd.to_numeric(dolar["Lebac"], errors='coerce')
dolar["NOCOM"] = pd.to_numeric(dolar["NOCOM"], errors='coerce')
dolar = dolar.fillna(0)

dolar["Pasivos"] = dolar["Base Monetaria"] + dolar["PasesPasivos"] - dolar["PasesActivos"] + dolar["Leliq"] + dolar["Lebac"] + dolar["LevidPesos"]

dolar["Dolar"] = dolar["Pasivos"] / dolar["Reservas"]

final = pd.merge(dolar, blue, on="Fecha", how="inner")[["Fecha", "Dolar", "Venta"]].copy().reset_index(drop=1)
final["Ratio"] = final["Dolar"]/final["Venta"]
final["Fecha"] = final["Fecha"].dt.strftime("%d/%m/%Y")
final.to_excel("TC_Convertibilidad.xlsx", index=False, sheet_name="Datos")

archivo_excel = openpyxl.load_workbook("TC_Convertibilidad.xlsx")
hoja_datos = archivo_excel['Datos']

datos_referencia_1 = Reference(hoja_datos,
                               min_col=2,
                               min_row=len(hoja_datos["B"])*.85,
                               max_row=len(hoja_datos["B"]))
serie_1 = Series(datos_referencia_1, title="Dólar Convertibilidad")

datos_referencia_2 = Reference(hoja_datos,
                               min_col=3,
                               min_row=len(hoja_datos["C"])*.85,
                               max_row=len(hoja_datos["C"]))
serie_2 = Series(datos_referencia_2, title="Dólar Blue")

datos_referencia_3 = Reference(hoja_datos,
                               min_col=4,
                               min_row=len(hoja_datos["D"])*.85,
                               max_row=len(hoja_datos["D"]))
serie_3 = Series(datos_referencia_3, title="Ratio")

grafico = LineChart()
grafico.title = "Dólar Convertibilidad"
grafico.x_axis.title = "Fecha"
grafico.y_axis.title = "Pesos"
grafico.height = 15
grafico.width = 30
grafico.legend.position = "b"
grafico.style = 4

grafico2 = LineChart()
grafico2.title = "Ratio Convertibilidad"
grafico2.x_axis.title = "Fecha"
grafico2.y_axis.title = "Ratio"
grafico2.height = 15
grafico2.width = 30
grafico2.legend.position = "b"
grafico2.style = 7

# Agregar las series a los gráficos.
grafico.series.append(serie_1)
grafico.series.append(serie_2)
grafico2.series.append(serie_3)

# Crear las categorías, eje X. Luego agregarlas a los gráficos.
categorias_referencia = Reference(hoja_datos,
                                  min_col=1,
                                  min_row=len(hoja_datos["A"])*.85,
                                  max_row=len(hoja_datos["A"]))
grafico.set_categories(categorias_referencia)
grafico2.set_categories(categorias_referencia)

# Agregar los gráficos a la hoja.
hoja_grafico = archivo_excel.create_sheet(title="Hoja del gráfico")
hoja_grafico.add_chart(grafico, "A1")
hoja_grafico.add_chart(grafico2, "A30")

# Guardar los cambios.
archivo_excel.save('TC_Convertibilidad_Excel.xlsx')
